import argparse, json, re
from pathlib import Path
from openpyxl import load_workbook
from data_contract import load_contract
from overrides import apply_meeting_day_moves, apply_physical_education_overrides, ensure_override_matches, load_day_moves, load_override_rules

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', type=Path)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--overrides', type=Path)
    args = parser.parse_args()
    ws = load_workbook(args.input, data_only=True).active
    weekdays = {3:1,4:2,5:3,6:4,7:5}; times = {}
    for row in range(2, ws.max_row + 1):
        label = str(ws.cell(row, 1).value or '')
        match = re.match(r'(\d+)', label)
        if match: times[row] = int(match.group(1))
    groups = {}
    for row, slot in times.items():
        for col, weekday in weekdays.items():
            value = ws.cell(row, col).value
            if not isinstance(value, str): continue
            for line in value.splitlines():
                line = line.strip()
                match = re.match(r'PE Group (\d+)\s+(.+?)\s+-\s+(.+)$', line)
                if not match: continue
                group, teacher, room = match.groups()
                groups.setdefault(group, []).append({'startWeek': 1, 'endWeek': 60, 'weekday': weekday, 'slot': slot, 'teachers': [teacher], 'room': room})
    override_path = args.overrides or Path('overrides') / '2026Autumn.yml'
    override_rules = load_override_rules(override_path)
    matched_overrides = apply_physical_education_overrides(groups, override_rules)
    ensure_override_matches(override_rules, matched_overrides, {'physicalEducation'})
    day_moves = load_day_moves(override_path)
    calendar_start = load_contract()['calendarStart']
    for group, meetings in groups.items(): groups[group] = apply_meeting_day_moves(meetings, day_moves, calendar_start)
    root = args.output / '2026Autumn'; root.mkdir(parents=True, exist_ok=True)
    for group, meetings in groups.items(): (root / f'pe-{group}.json').write_text(json.dumps({'groupId': group, 'meetings': meetings}, ensure_ascii=False) + '\n', encoding='utf-8')
if __name__ == '__main__': main()
